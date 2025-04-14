from openpyxl.reader.excel import load_workbook
import os
import pandas as pd
from openpyxl.utils import get_column_letter
from Searching.Searching import Searching

try:
    with Searching() as bot:
        bot.drushim()
        bot.intel()
        bot.amazon()
    for job in bot.jobs_found:
        for info in job:
            print(info)
        print('\n')
    #creating data frame
    df = pd.DataFrame(bot.jobs_found, columns=['title', 'location', 'date', 'link'])
    df.to_excel('output.xlsx', index=False)
    #opening workbook
    wb = load_workbook("output.xlsx")
    ws = wb.active
    #setting column widths
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 40
    #saving changes
    wb.save("output.xlsx")
    #open the excel file
    os.startfile("output.xlsx")



except Exception as e:
    if 'in Path' in str(e):
        print(
            'you are trying to run bot from command line\n'
            'please add to PATH your selenium drivers\n'
            'Windows: \n'
            'set PATH=%PATH%;C:*path-to-your-folder*\\\n'
            'Linux: \n'
            'PATH=$PATH:/path/toyour/folder/'
        )

    else:
        raise
